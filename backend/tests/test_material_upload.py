"""
Tayyor test / slayd fayllarini yuklash: matndan savollar va javob kaliti ajratiladi,
slaydlar tuziladi, asl fayl saqlanadi, talaba ro'yxatida file_url va items_count ko'rinadi.
Vaqtinchalik SQLite va vaqtinchalik uploads papkasi ishlatiladi.
"""
import io
import os
import tempfile

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.main import app
from app.database import Base, get_db
from app.core.security import create_access_token
from app.models.entities import User, TeacherSubject, SyllabusTopic
from app.services import material_upload as upsvc
from app.api import teacher as teacher_api


# ---------------- parser birlik testlari ----------------

def test_parse_yulduzcha_va_javob_qatori():
    text = """
1. Present Perfect qaysi gapda to'g'ri?
A) I have went there
B) I have gone there *
C) I has gone there
D) I gone there

2. "Book" so'zining ko'pligi?
A) Books
B) Bookes
C) Bookies
Javob: A

3) Savol uchinchi A) bir B) ikki C) uch
Answer: C
"""
    r = upsvc.parse_test_questions(text)
    qs = r["questions"]
    assert [q["correct_index"] for q in qs] == [1, 0, 2]
    assert qs[0]["options"][1] == "I have gone there"  # yulduzcha olib tashlanadi
    assert qs[2]["question"] == "Savol uchinchi" and qs[2]["options"] == ["bir", "ikki", "uch"]
    assert r["missing_keys"] == 0


def test_parse_kalit_bloki_va_kirill_harflar():
    text = """
1. Birinchi savol?
А) a
Б) b
В) c
2. Ikkinchi savol?
A. x
B. y
C. z
3. Kalitsiz savol?
A) p
B) q

Javoblar kaliti:
1-В, 2-A
"""
    r = upsvc.parse_test_questions(text)
    assert [q["correct_index"] for q in r["questions"]] == [2, 0, 0]
    assert r["missing_keys"] == 1  # 3-savolda kalit yo'q


def test_parse_bosh_matn():
    assert upsvc.parse_test_questions("Shunchaki matn, savol yo'q.")["questions"] == []


def test_slaydlar_txt_va_docx():
    txt = "1. Kirish\nBirinchi fikr\nIkkinchi fikr\n2. Asosiy qism\nUchinchi fikr\n"
    slides = upsvc.extract_slides(txt.encode("utf-8"), "slayd.txt")
    assert [s["title"] for s in slides] == ["1. Kirish", "2. Asosiy qism"]
    assert slides[0]["bullets"] == ["Birinchi fikr", "Ikkinchi fikr"]

    from docx import Document
    doc = Document()
    doc.add_heading("Mavzu A", level=1)
    doc.add_paragraph("A haqida")
    doc.add_heading("Mavzu B", level=1)
    doc.add_paragraph("B haqida")
    buf = io.BytesIO()
    doc.save(buf)
    slides = upsvc.extract_slides(buf.getvalue(), "slayd.docx")
    assert [s["title"] for s in slides] == ["Mavzu A", "Mavzu B"] and slides[1]["bullets"] == ["B haqida"]


def test_pptx_zip_tahlili():
    """python-pptx'siz: minimal .pptx (zip + slide XML) dan sarlavha, matn va notes olinadi."""
    import zipfile
    A = "http://schemas.openxmlformats.org/drawingml/2006/main"
    P = "http://schemas.openxmlformats.org/presentationml/2006/main"

    def sp(text, ph_type=None):
        ph = f'<p:nvPr><p:ph type="{ph_type}"/></p:nvPr>' if ph_type else "<p:nvPr/>"
        return f'<p:sp><p:nvSpPr><p:cNvPr id="1" name="x"/><p:cNvSpPr/>{ph}</p:nvSpPr><p:txBody><a:p><a:r><a:t>{text}</a:t></a:r></a:p></p:txBody></p:sp>'

    def slide_xml(*shapes):
        return f'<p:sld xmlns:a="{A}" xmlns:p="{P}"><p:cSld><p:spTree>{"".join(shapes)}</p:spTree></p:cSld></p:sld>'

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("ppt/slides/slide1.xml", slide_xml(sp("Kirish", "title"), sp("Birinchi tezis"), sp("Ikkinchi tezis")))
        z.writestr("ppt/slides/slide2.xml", slide_xml(sp("Xulosa", "ctrTitle"), sp("Yakuniy fikr")))
        z.writestr("ppt/notesSlides/notesSlide1.xml", slide_xml(sp("Spiker izohi"), sp("1")))
    slides = upsvc.extract_slides(buf.getvalue(), "taqdimot.pptx")
    assert [s["title"] for s in slides] == ["Kirish", "Xulosa"]
    assert slides[0]["bullets"] == ["Birinchi tezis", "Ikkinchi tezis"] and slides[0]["speaker_notes"] == "Spiker izohi"
    assert slides[1]["bullets"] == ["Yakuniy fikr"]
    assert "Kirish" in upsvc.extract_text(buf.getvalue(), "taqdimot.pptx")


# ---------------- endpoint testlari ----------------

@pytest.fixture()
def env(monkeypatch):
    tmp = tempfile.mkdtemp(prefix="tilora_upload_")
    engine = create_engine(f"sqlite:///{os.path.join(tmp, 't.db')}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(bind=engine)
    S = sessionmaker(bind=engine, autoflush=False, autocommit=False)

    def override_get_db():
        db = S()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    monkeypatch.setattr(teacher_api, "UPLOAD_MATERIAL_DIR", os.path.join(tmp, "materials"))

    db = S()
    db.add_all([
        User(id="t-1", email="t1@test.uz", password_hash="x", full_name="Ustoz 1", role="teacher", is_approved=True, subscription_tier="pro"),
        User(id="t-2", email="t2@test.uz", password_hash="x", full_name="Ustoz 2", role="teacher", is_approved=True),
        TeacherSubject(id="sub-a", teacher_id="t-1", name="Ingliz tili A", code="ENG-A", group_name="302-guruh"),
        SyllabusTopic(id="tp-1", subject_id="sub-a", week_number=1, title="Modal verbs", hours=2),
    ])
    db.commit()
    db.close()

    client = TestClient(app)
    h = lambda uid, role: {"Authorization": f"Bearer {create_access_token({'id': uid, 'role': role, 'subscription_tier': 'pro'})}"}
    yield client, h("t-1", "teacher"), h("t-2", "teacher"), tmp

    app.dependency_overrides.pop(get_db, None)
    engine.dispose()


TEST_TXT = b"1. Savol bir?\nA) a\nB) b *\n2. Savol ikki?\nA) x *\nB) y\nC) z\n"


def test_test_fayli_yuklanadi_va_talaba_ishlaydi(env):
    client, t1, t2, tmp = env
    r = client.post(
        "/api/teacher/subjects/sub-a/upload-material",
        data={"material_type": "test", "topic_title": "Modal verbs", "topic_id": "tp-1", "title": "Fayl testi"},
        files={"file": ("modal_test.txt", TEST_TXT, "text/plain")},
        headers=t1,
    )
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["items_count"] == 2 and d["file_only"] is False and d["file_url"].startswith("/uploads/materials/")
    assert os.path.exists(os.path.join(tmp, "materials", os.path.basename(d["file_url"])))
    mid = d["material_id"]

    # Talaba: javob kaliti sizmaydi, file_url ko'rinadi, test topshiriladi
    r = client.get("/api/student/materials?subject_id=sub-a")
    m = r.json()["materials"][0]
    assert m["items_count"] == 2 and m["content"]["file_url"] == d["file_url"]
    assert all("correct_index" not in q for q in m["content"]["questions"])
    r = client.post(f"/api/student/materials/{mid}/submit-test", json={"answers": {"q_1": 1, "q_2": 1}})
    assert r.status_code == 200 and r.json()["correct_count"] == 1

    # Boshqa o'qituvchi bu fanga yuklay olmaydi
    r = client.post("/api/teacher/subjects/sub-a/upload-material", data={"material_type": "test"},
                    files={"file": ("x.txt", TEST_TXT, "text/plain")}, headers=t2)
    assert r.status_code == 404


def test_savolsiz_fayl_rad_etiladi_file_rejimi_qabul_qilinadi(env):
    client, t1, _, tmp = env
    body = b"Bu oddiy matn, test emas."
    r = client.post("/api/teacher/subjects/sub-a/upload-material", data={"material_type": "test"},
                    files={"file": ("oddiy.txt", body, "text/plain")}, headers=t1)
    assert r.status_code == 400 and "ajratilmadi" in r.json()["detail"]

    r = client.post("/api/teacher/subjects/sub-a/upload-material", data={"material_type": "test", "mode": "file"},
                    files={"file": ("oddiy.txt", body, "text/plain")}, headers=t1)
    assert r.status_code == 200 and r.json()["file_only"] is True and r.json()["items_count"] == 0

    # Noto'g'ri kengaytma
    r = client.post("/api/teacher/subjects/sub-a/upload-material", data={"material_type": "test"},
                    files={"file": ("x.exe", b"abc", "application/octet-stream")}, headers=t1)
    assert r.status_code == 400


def test_slayd_fayli_yuklanadi(env):
    client, t1, _, tmp = env
    body = "1. Kirish\nA fikr\nB fikr\n2. Xulosa\nC fikr\n".encode("utf-8")
    r = client.post("/api/teacher/subjects/sub-a/upload-material", data={"material_type": "presentation"},
                    files={"file": ("slayd.txt", body, "text/plain")}, headers=t1)
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["items_count"] == 2 and d["data"]["slides"][1]["title"] == "2. Xulosa"
    r = client.get("/api/student/materials?material_type=presentation")
    m = r.json()["materials"][0]
    assert m["items_count"] == 2 and m["content"]["file_name"] == "slayd.txt"

"""
Sillabus va o'quv dasturi hujjatlarini (.docx, .pdf, .txt) intellektual tahlil qilish va
mavzularni avtomatik ajratib olish xizmati (Syllabus Document Parser).
O'zbekiston Oliy ta'lim muassasalari standart o'quv dasturlari (HEMIS, Vazirlik namunaviy dasturlari,
CamScanner skanerlangan PDF va Word formatlari)ga to'liq moslashtirilgan.
"""

import io
import os
import re
import sys
import tempfile
import subprocess
from typing import Dict, Any, List, Optional, Tuple

try:
    import docx
except ImportError:
    docx = None

try:
    import fitz
except ImportError:
    try:
        import pymupdf as fitz
    except ImportError:
        fitz = None


def normalize_apostrophes(text: str) -> str:
    """O'zbek tilidagi barcha turli apostrof belgilarni bir xil ko'rinishga keltiradi."""
    if not text:
        return ""
    return re.sub(r"[\u02bc\u2018\u2019`ʻʼ]", "'", text)


def clean_subject_name(raw_name: str) -> str:
    """Fan nomini tozalash va chiroyli holga keltirish."""
    name = normalize_apostrophes(raw_name).strip()
    name = re.sub(r"^(?:fanning\s+nomi|fan\s+nomi|subject\s+name|fan)\s*[:\-–—]?\s*", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\s+(?:fani|fani\s+bo'yicha\s+o'quv\s+dasturi|fanining\s+o'quv\s+dasturi).*$", "", name, flags=re.IGNORECASE)
    name = name.strip(" .:-–—\"'«»“”\t\r\n")
    if name.isupper() and len(name) > 3:
        parts = [p.capitalize() for p in name.split()]
        name = " ".join(parts)
        name = re.sub(r"\bO'([a-z])", lambda m: "O'" + m.group(1).lower(), name)
        name = re.sub(r"\bG'([a-z])", lambda m: "G'" + m.group(1).lower(), name)
    return name


def deduce_subject_name_from_filename(filename: str) -> str:
    """Fayl nomidan fan nomini ajratib olish (fallback)."""
    base = os.path.basename(filename)
    base = os.path.splitext(base)[0]
    base = re.sub(r'^[0-9_\-\.\s]+', '', base)
    base = re.sub(r'[\-_]?(?:fan[i_]?[\-_]?dasturi|dastur(?:i)?|sillabus|syllabus).*$', '', base, flags=re.IGNORECASE)
    base = re.sub(r'[\-_]?[A-Za-z0-9\u02bc\u2018\u2019\']{4,12}$', '', base)
    base = base.replace('_', ' ').strip()
    return clean_subject_name(base)


def parse_syllabus_document(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """
    Yuklangan fayl (.docx, .pdf, .txt) baytlaridan fan ma'lumotlari va
    barcha mavzularni (Ma'ruza, Seminar, Mustaqil ta'lim) ajratib oladi.
    """
    lower_name = filename.lower()
    is_pdf = lower_name.endswith(".pdf") or file_bytes.startswith(b"%PDF-")
    is_docx = lower_name.endswith(".docx") or file_bytes.startswith(b"PK\x03\x04")
    
    if is_docx:
        return _parse_docx_bytes(file_bytes, filename)
    elif is_pdf:
        return _parse_pdf_bytes(file_bytes, filename)
    else:
        return _parse_txt_bytes(file_bytes, filename)


def _extract_metadata_from_blocks(blocks: List[Tuple[str, bool]], filename: str) -> Tuple[str, str, str, str]:
    """Hujjat bloklaridan fan nomi, kodi, semestr va tavsifni intellektual ajratib olish."""
    subject_name = ""
    subject_code = ""
    semester = ""
    description = ""

    lines = [normalize_apostrophes(b[0]).strip() for b in blocks if b[0].strip()]
    full_text = "\n".join(lines)

    # 1. FAN NOMI:
    # 1.a) Sarlavha qatorlaridan tekshirish: "[FAN NOMI] FANI BO'YICHA O'QUV DASTURI"
    re_title_marker = re.compile(r"fani(?:ng)?\s+(?:bo['’`]?yic[hi]{1,2}a\s+)?o['’`]?quv\s+dasturi|fanidan\s+sillabus", re.IGNORECASE)
    for i, line in enumerate(lines[:40]):
        if re_title_marker.search(line):
            if i > 0 and len(lines[i-1]) > 3 and not any(k in lines[i-1].upper() for k in ["RESPUBLIKASI", "VAZIRLIGI", "UNIVERSITETI", "TASDIQLAYMAN"]):
                subject_name = clean_subject_name(lines[i-1])
                break
        m_title = re.search(r"([A-Za-z\u0400-\u04FF\s']{4,60})\s+fani(?:ng)?\s+(?:bo['’`]?yic[hi]{1,2}a\s+)?o['’`]?quv\s+dasturi", line, re.IGNORECASE)
        if m_title:
            cand = clean_subject_name(m_title.group(1))
            if not any(k in cand.upper() for k in ["RESPUBLIKASI", "VAZIRLIGI", "UNIVERSITETI"]):
                subject_name = cand
                break

    # 1.b) Jadvallardan "Fanning nomi" izlash
    if not subject_name:
        for i, line in enumerate(lines[:60]):
            m_fn = re.search(r"(?:Fanning\s+nomi|Fan\s+nomi)\s*[:\-–—]?\s*(.+)?", line, re.IGNORECASE)
            if m_fn:
                val = m_fn.group(1) if m_fn.group(1) else ""
                val = val.strip()
                if val and len(val) > 3 and not any(k in val.lower() for k in ["auditoriya", "soat", "kredit", "turi", "fanning nomi"]):
                    subject_name = clean_subject_name(val)
                    break
                elif i + 1 < len(lines):
                    next_val = lines[i+1].strip()
                    if next_val and len(next_val) > 3 and not any(k in next_val.lower() for k in ["auditoriya", "soat", "kredit", "fanning", "fan/"]):
                        subject_name = clean_subject_name(next_val)
                        break

    # 1.c) Fayl nomidan aniqlash (fallback)
    if not subject_name or len(subject_name) < 3 or any(bad in subject_name.lower() for bad in ["auditoriya", "mashg'ulot", "ta'lim"]):
        subject_name = deduce_subject_name_from_filename(filename)

    # 2. FAN KODI:
    m_code = re.search(r"(?:Fan(?:/modul)?|Modul|Fnn/[a-z]+)\s+kodi\s*[:\-–—]?\s*([A-Za-z0-9\-\u02bc\u2018\u2019']+)", full_text, re.IGNORECASE)
    if m_code:
        subject_code = normalize_apostrophes(m_code.group(1)).strip(" -:")
    else:
        m_code2 = re.search(r"\b([A-Z]{2,6}[0-9]{3,5}|[A-Z]{1,4}'?[A-Z]{2,4}[0-9]{3,5})\b", full_text)
        if m_code2:
            subject_code = m_code2.group(1).strip()
        else:
            m_fn_code = re.search(r'[\-_]([A-Za-z0-9\u02bc\u2018\u2019\']{4,12})(?:[\-_]fan|\.|\Z)', filename)
            if m_fn_code:
                subject_code = normalize_apostrophes(m_fn_code.group(1)).strip()

    # 3. SEMESTR:
    m_sem = re.search(r"Semestr\s*[:\-–—]?\s*(\d+)", full_text, re.IGNORECASE)
    if m_sem:
        semester = f"{m_sem.group(1)}-semestr"
    else:
        m_sem2 = re.search(r"(\d+)[-–—\s]?semestr", full_text, re.IGNORECASE)
        if m_sem2:
            semester = f"{m_sem2.group(1)}-semestr"

    # 4. TAVSIF / MAQSAD:
    m_desc = re.search(r"(?:Fann?i(?:ng)?\s+o'qitishning\s+maqsadi|Fann?ing\s+maqsadi|Fan\s+maqsadi)\s*[:\-–—]?\s*([^\n\r]+(?:\n[^\n\r]+){0,3})", full_text, re.IGNORECASE)
    if m_desc:
        clean_d = m_desc.group(1).strip()
        description = re.sub(r"\s+", " ", clean_d)

    return subject_name, subject_code, semester or "1-semestr", description


def _extract_topics_from_blocks(blocks: List[Tuple[str, bool]], filename: str = "") -> Dict[str, Any]:
    sec = "none"
    topics_by_sec = {
        "lecture": [],
        "seminar": [],
        "independent": []
    }
    
    re_lecture_header = re.compile(r"(?:asosiy\s+nazariy\s+qism|nnznriy\s+qism|ma[’'‘`\.\*]?ruza(?:\s*mashg[’'‘`]ulotlari)?|mn•ruza|ma[’'‘`]ruzalar|lectures?)", re.IGNORECASE)
    re_seminar_header = re.compile(r"(?:seminar(?:\s*mashg[’'‘`]ulotlar)?|amaliy(?:\s*mashg[’'‘`]ulotlar)?|seminar\s+va\s+amaliy|laboratoriya|seminars?|practicals?)", re.IGNORECASE)
    re_indep_header = re.compile(r"(?:mustaqil\s*ta[’'‘`]?lim|mustaqil\s*ishlar|mustaqil\s*ta[’'‘`]?lim\s+topshiriqlari|independent\s*study)", re.IGNORECASE)
    
    re_stop_header = re.compile(
        r"(?:fan\s+o[’'‘`]?qitilishining\s+natijalari|shakllanadigan\s+kompetensiyalar|ta[’'‘`]?lim\s+texnologiyalari|"
        r"kreditlarni\s+olish\s+uchun\s+talablar|rahbariy\s+adabiyotlar|asosiy\s+(?:va\s+qo[’'‘`]?shimcha\s+)?adabiyotlar|"
        r"qo[’'‘`]?shimcha\s+adabiyotlar|adabiyotlar\s+ro[’'‘`]?yxati|adabiyotlar:|axborot\s+manbalari|foydalanilgan\s+adabiyotlar|"
        r"baholash\s+mezon|taqrizchilar|fan/modul\s+uchun\s+ma[’'‘`]?sullar|kengashining\s+\d+|nazorat\s+savollari)",
        re.IGNORECASE
    )

    re_mavzu_line = re.compile(
        r"^(\d+|[tToO])\s*[-–—,\*\.rR]?\s*(?:mavzu|тема|topic|lecture|dars|[a-zA-Z\u00c0-\u024f]{1,4}[nv][nvz7u]{1,3})\b[\.:]?\s*(.+)",
        re.IGNORECASE
    )
    re_numbered_line = re.compile(r"^(\d{1,2})[\.\)]\s*(.+)")

    for raw_text, is_bold in blocks:
        t_clean = normalize_apostrophes(raw_text).strip()
        if not t_clean or len(t_clean) < 2:
            continue
            
        t_lower = t_clean.lower()

        if any(bad in t_lower for bad in [
            "www.", "http://", "https://", "camscanner", "toshkent:", "nashriyot", "darslik",
            "fan/modul kodi", "fan/modul turi", "haftadagi dars soatlari", "auditoriya mashg'ulotlari",
            "umumiy o'quv soati", "kafedrasi mudiri"
        ]):
            continue

        if re_stop_header.search(t_clean):
            sec = "other"
            continue
        elif re_seminar_header.search(t_clean) and ("tavsiya" in t_lower or "mavzular" in t_lower or "mashg'ulot" in t_lower or is_bold or len(t_clean) < 80):
            sec = "seminar"
            continue
        elif re_indep_header.search(t_clean) and ("tavsiya" in t_lower or "mavzular" in t_lower or "topshiriq" in t_lower or is_bold or len(t_clean) < 80):
            sec = "independent"
            continue
        elif re_lecture_header.search(t_clean) and ("asosiy" in t_lower or "mavzular" in t_lower or "mashg'ulot" in t_lower or is_bold or len(t_clean) < 80):
            sec = "lecture"
            continue

        if sec == "other":
            continue

        target_sec = sec if sec in topics_by_sec else "lecture"
        curr_list = topics_by_sec[target_sec]

        m_mavzu = re_mavzu_line.match(t_clean)
        m_num = re_numbered_line.match(t_clean)

        if m_mavzu:
            num_str = m_mavzu.group(1)
            if num_str.isdigit():
                num_val = int(num_str)
            elif num_str.lower() in ('t', 'i', 'l', '|'):
                num_val = 1
            else:
                num_val = len(curr_list) + 1

            title = m_mavzu.group(2).strip()
            
            if sec == "none":
                sec = "lecture"
                curr_list = topics_by_sec["lecture"]

            curr_list.append({
                "code": f"M{num_val}",
                "num": num_val,
                "title": title,
                "subtopics": [],
                "hours": 2,
                "is_bold": is_bold
            })
            
        elif m_num and (sec in ["seminar", "independent"] or is_bold or (sec == "lecture" and len(m_num.group(2)) > 10)):
            num_val = int(m_num.group(1))
            title = m_num.group(2).strip()
            
            is_author = any(k in title for k in ["Mirziyoyev", "Murtazayeva", "Sagdullayev", "Eshov", "Qonun", "Farmon", "Qaror"])
            if len(title) >= 5 and not is_author:
                curr_list.append({
                    "code": f"M{num_val}",
                    "num": num_val,
                    "title": title,
                    "subtopics": [],
                    "hours": 2,
                    "is_bold": is_bold
                })
        else:
            if curr_list:
                last = curr_list[-1]
                should_join = False
                
                if last["is_bold"] and is_bold and not re.match(r"^(\d+|[A-ZIVX]+[\.\)])", t_clean):
                    should_join = True
                elif not is_bold and not last["title"].rstrip().endswith((".", "!", "?")) and len(last["title"]) < 70 and t_clean[0].islower():
                    should_join = True
                    
                if should_join:
                    last["title"] += " " + t_clean
                else:
                    if len(last["subtopics"]) < 10 and not any(bad in t_lower for bad in ["buxoro 20", "rektor", "tasdiqlayman"]):
                        last["subtopics"].append(t_clean)

    subject_name, subject_code, semester, description = _extract_metadata_from_blocks(blocks, filename)

    lecture_topics = []
    for idx, t in enumerate(topics_by_sec["lecture"], 1):
        clean_t = re.sub(r"^(?:\d+[-–—,\*\.rR]?\s*(?:mavzu|[a-zA-Z\u00c0-\u024f]{1,4}[nv][nvz7u]{1,3})[\.:]?|\d+[\.\)])\s*", "", t["title"]).strip()
        lecture_topics.append({
            "category": "lecture",
            "code": f"M{idx}",
            "title": clean_t,
            "subtopics": " ".join(t["subtopics"])[:500],
            "hours": t.get("hours", 2)
        })

    seminar_topics = []
    for idx, t in enumerate(topics_by_sec["seminar"], 1):
        clean_t = re.sub(r"^(?:\d+[-–—,\*\.rR]?\s*(?:mavzu|[a-zA-Z\u00c0-\u024f]{1,4}[nv][nvz7u]{1,3})[\.:]?|\d+[\.\)])\s*", "", t["title"]).strip()
        seminar_topics.append({
            "category": "seminar",
            "code": f"S{idx}",
            "title": clean_t,
            "subtopics": " ".join(t["subtopics"])[:500],
            "hours": t.get("hours", 2)
        })

    independent_topics = []
    for idx, t in enumerate(topics_by_sec["independent"], 1):
        clean_t = re.sub(r"^(?:\d+[-–—,\*\.rR]?\s*(?:mavzu|[a-zA-Z\u00c0-\u024f]{1,4}[nv][nvz7u]{1,3})[\.:]?|\d+[\.\)])\s*", "", t["title"]).strip()
        independent_topics.append({
            "category": "independent",
            "code": f"I{idx}",
            "title": clean_t,
            "subtopics": " ".join(t["subtopics"])[:500],
            "hours": t.get("hours", 2)
        })

    all_topics = lecture_topics + seminar_topics + independent_topics
    if not lecture_topics and all_topics:
        lecture_topics = all_topics

    def format_list(top_list):
        return "\n".join(f"{i}. {t['title']}" for i, t in enumerate(top_list, 1))

    return {
        "subject_name": subject_name,
        "subject_code": subject_code,
        "semester": semester or "1-semestr",
        "description": description,
        "lecture_topics": lecture_topics,
        "seminar_topics": seminar_topics,
        "independent_topics": independent_topics,
        "all_topics": all_topics,
        "formatted_text": format_list(lecture_topics) if lecture_topics else format_list(all_topics),
        "lecture_text": format_list(lecture_topics),
        "seminar_text": format_list(seminar_topics),
        "independent_text": format_list(independent_topics)
    }


def _parse_docx_bytes(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    if docx is None:
        raise RuntimeError("python-docx kutubxonasi o'rnatilmagan.")
        
    stream = io.BytesIO(file_bytes)
    doc = docx.Document(stream)
    
    # 1. Jadval ichida har bir qatori mavzu bo'lgan (M1, M2... yoki S1, S2...) jadval bormi?
    dedicated_topic_table = None
    for t in doc.tables:
        has_codes = 0
        for r in t.rows:
            distinct = []
            seen = []
            for c in r.cells:
                if c._tc not in seen:
                    seen.append(c._tc)
                    distinct.append(c.text.strip())
            if distinct and re.match(r'^(?:[MSLA]|MT)\d+$', distinct[0], re.IGNORECASE):
                has_codes += 1
        if has_codes >= 5:
            dedicated_topic_table = t
            break

    if dedicated_topic_table:
        # 1-format: Standart qatorli jadval (Dastur Boshlang'ich ta'lim kabi)
        extracted = {
            "subject_name": "",
            "subject_code": "",
            "semester": "",
            "description": "",
            "lecture_topics": [],
            "seminar_topics": [],
            "independent_topics": [],
            "all_topics": [],
            "formatted_text": "",
            "lecture_text": "",
            "seminar_text": "",
            "independent_text": ""
        }

        # Barcha jadvallardan metadatalarni olish
        for table in doc.tables:
            for row in table.rows:
                seen = []
                distinct = []
                for c in row.cells:
                    if c._tc not in seen:
                        seen.append(c._tc)
                        distinct.append(c.text.strip())
                row_text = " ".join(distinct)
                if "Fanning nomi" in row_text or "Fan nomi" in row_text:
                    for cell in distinct:
                        m = re.search(r"Fann?ing\s+nomi\s*[:\-]?\s*(.+)", cell, re.IGNORECASE)
                        if m:
                            val = m.group(1).strip()
                            if val and len(val) > 2 and "fanning nomi" not in val.lower():
                                extracted["subject_name"] = clean_subject_name(val)
                if "Fan/modul kodi" in row_text or "Fan kodi" in row_text or "Modul kodi" in row_text:
                    for cell in distinct:
                        m = re.search(r"(?:Fan(?:/modul)?|Modul)\s+kodi\s*[:\-]?\s*([A-Za-z0-9\-]+)", cell, re.IGNORECASE)
                        if m:
                            extracted["subject_code"] = m.group(1).strip()
                if "Semestr" in row_text:
                    for cell in distinct:
                        m = re.search(r"Semestr\s*[:\-]?\s*([0-9,\s\-]+)", cell, re.IGNORECASE)
                        if m:
                            extracted["semester"] = f"{m.group(1).strip()}-semestr"
                if "Fan maqsadi:" in row_text or "Fanning maqsadi:" in row_text:
                    for cell in distinct:
                        if "Fan maqsadi:" in cell or "Fanning maqsadi:" in cell:
                            clean_desc = re.sub(r"Fann?ing\s+maqsadi:\s*", "", cell, flags=re.IGNORECASE).strip()
                            extracted["description"] = clean_desc

        if not extracted["subject_name"]:
            extracted["subject_name"] = deduce_subject_name_from_filename(filename)

        curr_category = "lecture"
        for row in dedicated_topic_table.rows:
            seen = []
            distinct = []
            for c in row.cells:
                if c._tc not in seen:
                    seen.append(c._tc)
                    distinct.append(c.text.replace("\n", " ").strip())

            if not distinct:
                continue

            row_txt = " ".join(distinct).upper()
            if "MA’RUZA" in row_txt or "MA'RUZA" in row_txt or "LECTURE" in row_txt:
                curr_category = "lecture"
                continue
            elif "SEMINAR" in row_txt or "AMALIY" in row_txt or "PRACTICAL" in row_txt:
                curr_category = "seminar"
                continue
            elif "MUSTAQIL" in row_txt or "INDEPENDENT" in row_txt:
                curr_category = "independent"
                continue

            c0 = distinct[0].strip()
            if len(distinct) >= 2 and re.match(r'^(?:[MSLA]|MT)?\d+$', c0, re.IGNORECASE):
                topic_title = distinct[1].strip()
                if len(topic_title) < 4 or any(bad in topic_title.lower() for bad in ["fan ma'lumotlari", "fanning nomi", "baholash"]):
                    continue

                hours = 2
                if len(distinct) >= 3:
                    hm = re.search(r'(\d+)', distinct[2])
                    if hm:
                        hours = int(hm.group(1))

                parts = [p.strip() for p in re.split(r'(?<=[.!?])\s+', topic_title) if p.strip()]
                main_title = parts[0] if parts else topic_title
                subtopics = " ".join(parts[1:]) if len(parts) > 1 else ""

                item = {
                    "category": curr_category,
                    "code": c0,
                    "title": main_title,
                    "subtopics": subtopics,
                    "hours": hours
                }

                if curr_category == "lecture":
                    extracted["lecture_topics"].append(item)
                elif curr_category == "seminar":
                    extracted["seminar_topics"].append(item)
                elif curr_category == "independent":
                    extracted["independent_topics"].append(item)
                extracted["all_topics"].append(item)

        def format_list(top_list):
            return "\n".join(f"{i}. {t['title']}" for i, t in enumerate(top_list, 1))

        chosen = extracted["lecture_topics"] if extracted["lecture_topics"] else extracted["all_topics"]
        extracted["formatted_text"] = format_list(chosen)
        extracted["lecture_text"] = format_list(extracted["lecture_topics"])
        extracted["seminar_text"] = format_list(extracted["seminar_topics"])
        extracted["independent_text"] = format_list(extracted["independent_topics"])
        return extracted

    # 2-format: Katak ichidagi o'quv dasturi yoki matnli o'quv dasturi (O'zbekistonning eng yangi tarixi kabi)
    seen_tc = set()
    blocks = []
    
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            blocks.append((t, any(r.bold for r in p.runs)))
            
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in seen_tc:
                    continue
                seen_tc.add(cell._tc)
                for p in cell.paragraphs:
                    t = p.text.strip()
                    if t:
                        blocks.append((t, any(r.bold for r in p.runs)))

    return _extract_topics_from_blocks(blocks, filename)


def _parse_pdf_bytes(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    if fitz is None:
        raise RuntimeError("PyMuPDF (fitz) kutubxonasi o'rnatilmagan.")
        
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    
    extracted_text_pages = []
    total_chars = 0
    for page in doc:
        ptxt = page.get_text()
        extracted_text_pages.append(ptxt)
        total_chars += len(ptxt.strip())
        
    blocks = []
    
    if total_chars > 200:
        for ptxt in extracted_text_pages:
            for line in ptxt.splitlines():
                t = line.strip()
                if t:
                    blocks.append((t, False))
    else:
        temp_imgs = []
        try:
            max_pages = min(14, len(doc))
            for i in range(max_pages):
                pix = doc[i].get_pixmap(dpi=130)
                fd, path = tempfile.mkstemp(suffix=f"_ocr_p{i}.png")
                os.close(fd)
                pix.save(path)
                temp_imgs.append(path)

            ps1_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "ocr_helper.ps1")
            if not os.path.exists(ps1_path):
                ps1_path = r"C:\Users\HP\Desktop\TILORA\backend\app\services\ocr_helper.ps1"

            cmd = ["powershell", "-ExecutionPolicy", "Bypass", "-File", ps1_path] + temp_imgs
            res = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8")
            
            for part in res.stdout.split("=== PAGE_START ===")[1:]:
                sub = part.split("=== PAGE_END ===")[0].strip()
                for line in sub.splitlines():
                    t = line.strip()
                    if t:
                        blocks.append((t, False))
        finally:
            for p in temp_imgs:
                if os.path.exists(p):
                    try:
                        os.remove(p)
                    except Exception:
                        pass

    return _extract_topics_from_blocks(blocks, filename)


def _parse_txt_bytes(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    text = ""
    for enc in ["utf-8", "windows-1251", "cp1254", "latin-1"]:
        try:
            text = file_bytes.decode(enc)
            break
        except UnicodeDecodeError:
            continue
            
    blocks = [(l.strip(), False) for l in text.splitlines() if l.strip()]
    return _extract_topics_from_blocks(blocks, filename)
